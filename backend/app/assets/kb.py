"""KB ingestion (Pass 3): parse → chunk (ACTUALLY applied — carried lesson) →
embed via the model adapter when one is configured. No provider → chunks are
stored unembedded and the source is honestly marked keyword-only; nothing is
fabricated.
"""
from __future__ import annotations

import csv
import io
import re
import time

from sqlalchemy import delete
from sqlalchemy.orm import Session

from ..adapters import files as file_store
from ..adapters.models import ModelCallError, ModelUnavailable, get_model_adapter
from ..config import settings
from ..models import KbChunk, KnowledgeSource, utcnow

# Transient embedding failures (quota, network blip, provider 5xx) are worth a
# second try; a bounded budget keeps an upload request from hanging. Exhausting
# it changes nothing about the outcome — the existing degradation path runs.
EMBED_MAX_ATTEMPTS = 3
EMBED_BACKOFF_SECONDS = (0.5, 1.5)


def _decode(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _parse_docx(data: bytes) -> list[tuple[str, str]]:
    """Paragraphs in document order, then each table separately so a table's
    rows stay together rather than being interleaved into prose."""
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ValueError(f"DOCX support requires python-docx: {exc}") from exc
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f"not a readable .docx file: {exc}") from exc

    units: list[tuple[str, str]] = []
    paragraphs: list[str] = []
    table_no = 0
    # walk the body so paragraph/table ordering follows the document, which
    # document.paragraphs / document.tables alone would lose
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, document).text.strip()
            if text:
                paragraphs.append(text)
        elif tag == "tbl":
            table_no += 1
            rows: list[str] = []
            for row in Table(child, document).rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                units.append(("\n".join(rows), f"table {table_no}"))
    if paragraphs:
        units.insert(0, ("\n".join(paragraphs), "document"))
    return units


def _parse_html(data: bytes) -> list[tuple[str, str]]:
    """Readable text from an UPLOADED html file. No network access."""
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ValueError(f"HTML support requires beautifulsoup4: {exc}") from exc
    soup = BeautifulSoup(_decode(data), "html.parser")
    for element in soup(["script", "style", "noscript", "template"]):
        element.decompose()
    text = re.sub(r"\n{3,}", "\n\n", soup.get_text(separator="\n")).strip()
    return [(text, "document")] if text else []


def _parse_csv(data: bytes) -> list[tuple[str, str]]:
    """Header-aware rows as `col: value` pairs — stdlib only, no dataframe
    semantics and no structured-query behaviour."""
    reader = csv.reader(io.StringIO(_decode(data)))
    try:
        rows = [r for r in reader if any((c or "").strip() for c in r)]
    except csv.Error as exc:
        raise ValueError(f"not a readable .csv file: {exc}") from exc
    if not rows:
        return []
    header, *body = rows
    if not body:
        return [(" | ".join(c.strip() for c in header), "document")]
    lines = [
        " | ".join(f"{(header[i] if i < len(header) else f'col{i + 1}').strip()}: {(cell or '').strip()}"
                   for i, cell in enumerate(row) if (cell or "").strip())
        for row in body
    ]
    return [("\n".join(line for line in lines if line), "document")]


def _parse_xlsx(data: bytes) -> list[tuple[str, str]]:
    """One unit per worksheet. data_only=True reads cached values, so formulas
    are never evaluated and macros are never executed."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ValueError(f"XLSX support requires openpyxl: {exc}") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"not a readable .xlsx file: {exc}") from exc

    units: list[tuple[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            lines: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    lines.append(" | ".join(cells))
            if lines:
                units.append(("\n".join(lines), f"sheet {sheet.title}"))
    finally:
        workbook.close()
    return units


def extract_units(data: bytes, filename: str) -> list[tuple[str, str]]:
    """→ [(text, location_label)]. PDF = per page; XLSX = per sheet; DOCX =
    body + one unit per table; HTML/CSV/MD/TXT = whole file."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        units: list[tuple[str, str]] = []
        for i, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                units.append((text, f"page {i}"))
        return units
    if lower.endswith(".docx"):
        return _parse_docx(data)
    if lower.endswith((".html", ".htm")):
        return _parse_html(data)
    if lower.endswith(".csv"):
        return _parse_csv(data)
    if lower.endswith(".xlsx"):
        return _parse_xlsx(data)
    text = _decode(data).strip()
    return [(text, "document")] if text else []


def embed_with_retry(adapter, texts: list[str]) -> list[list[float]]:
    """Retry the EXISTING adapter call on transient failures only.

    The provider taxonomy already distinguishes these: ModelUnavailable means
    "not configured / SDK missing", which retrying cannot fix, so it is raised
    immediately. ModelCallError covers quota, network, and provider 5xx — worth
    one or two more attempts. When the budget is exhausted the original
    exception propagates, so the caller's degradation path is unchanged.
    """
    last: ModelCallError | None = None
    for attempt in range(EMBED_MAX_ATTEMPTS):
        try:
            return adapter.embed(texts)
        except ModelUnavailable:
            raise  # permanent — no amount of retrying configures a provider
        except ModelCallError as exc:
            last = exc
            if attempt < EMBED_MAX_ATTEMPTS - 1:
                time.sleep(EMBED_BACKOFF_SECONDS[min(attempt, len(EMBED_BACKOFF_SECONDS) - 1)])
    raise last  # type: ignore[misc]  # unreachable unless the loop ran


def chunk_units(units: list[tuple[str, str]], size: int, overlap: int) -> list[tuple[str, dict]]:
    """Fixed-size + overlap chunking within each unit. → [(text, meta)]."""
    if size <= 0:
        raise ValueError("chunk size must be positive")
    overlap = min(overlap, size - 1) if size > 1 else 0
    out: list[tuple[str, dict]] = []
    for text, location in units:
        start = 0
        while start < len(text):
            end = min(start + size, len(text))
            piece = text[start:end].strip()
            if piece:
                out.append((piece, {"location": f"{location}, chars {start}-{end}"}))
            if end == len(text):
                break
            start = end - overlap
    return out


def ingest(db: Session, source: KnowledgeSource, chunk_size: int = 800, chunk_overlap: int = 120) -> None:
    """(Re)ingest a source: replaces its chunks. Chunks are drafts of derived
    data, not governed records — replacement is legitimate here."""
    source.status = "ingesting"
    db.flush()
    try:
        data = file_store.read_bytes(source.file_path)
        units = extract_units(data, source.filename)
        if not units:
            raise ValueError("no extractable text found in the file")
        chunks = chunk_units(units, chunk_size, chunk_overlap)
        if not chunks:
            raise ValueError("chunking produced no chunks")

        db.execute(delete(KbChunk).where(KbChunk.source_id == source.id))

        # embedding model has its own catalog entry, so its own credential
        adapter = get_model_adapter(db, settings.gemini_embedding_model)
        embeddings: list[list[float]] | None = None
        if adapter is not None:
            try:
                embeddings = embed_with_retry(adapter, [text for text, _ in chunks])
            except (ModelUnavailable, ModelCallError) as exc:
                embeddings = None
                source.error = f"embedding unavailable: {exc}"  # honest, non-fatal

        for i, (text, meta) in enumerate(chunks):
            db.add(KbChunk(
                source_id=source.id, ord=i, text=text, meta=meta,
                embedding=embeddings[i] if embeddings else None,
            ))
        source.chunk_count = len(chunks)
        source.embedded = embeddings is not None
        source.status = "ready"
        if embeddings is not None:
            source.error = None
    except Exception as exc:
        source.status = "failed"
        source.error = str(exc)[:500]
    db.flush()
