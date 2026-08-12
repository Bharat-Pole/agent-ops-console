"""Document parsers (DOCX/HTML/CSV/XLSX) and bounded embedding retry.

Both additions are additive: existing PDF/MD/TXT extraction, chunking, and the
embedding degradation outcome are asserted unchanged alongside the new work.
"""
from __future__ import annotations

import io

import pytest
from conftest import login

from app.adapters import models as adapters
from app.adapters.models import ModelCallError, ModelUnavailable
from app.assets import kb


@pytest.fixture(autouse=True)
def _clear_adapter():
    yield
    adapters.set_adapter_override(None)


# ---- builders for small in-memory documents ---------------------------------

def _docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document
    document = Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---- DOCX -------------------------------------------------------------------

def test_docx_paragraphs_and_tables():
    data = _docx_bytes(
        ["Escalation policy", "Call the NOC lead within 15 minutes."],
        table=[["Severity", "Response"], ["P1", "15 minutes"]],
    )
    units = kb.extract_units(data, "policy.docx")
    labels = [label for _, label in units]
    assert "document" in labels and "table 1" in labels

    body = next(text for text, label in units if label == "document")
    assert "Escalation policy" in body
    assert "Call the NOC lead" in body

    table = next(text for text, label in units if label == "table 1")
    assert "Severity | Response" in table
    assert "P1 | 15 minutes" in table


def test_docx_empty_document_yields_no_units():
    assert kb.extract_units(_docx_bytes([]), "empty.docx") == []


def test_malformed_docx_fails_cleanly():
    with pytest.raises(ValueError, match="not a readable .docx"):
        kb.extract_units(b"this is definitely not a docx archive", "broken.docx")


# ---- HTML -------------------------------------------------------------------

def test_html_extracts_visible_text_and_drops_script_style():
    html = b"""<html><head><style>.x{color:red}</style>
    <script>var secret='do-not-index';</script></head>
    <body><h1>Runbook</h1><p>Check the router logs.</p>
    <noscript>enable js</noscript></body></html>"""
    units = kb.extract_units(html, "page.html")
    assert len(units) == 1 and units[0][1] == "document"
    text = units[0][0]
    assert "Runbook" in text and "Check the router logs." in text
    assert "do-not-index" not in text  # script contents excluded
    assert "color:red" not in text     # style contents excluded
    assert "enable js" not in text


def test_htm_extension_also_supported():
    units = kb.extract_units(b"<html><body><p>hello</p></body></html>", "page.htm")
    assert units and "hello" in units[0][0]


# ---- CSV --------------------------------------------------------------------

def test_csv_rows_are_header_aware_and_deterministic():
    data = b"incident,severity,owner\nINC-1,P1,noc\nINC-2,P3,support\n"
    units = kb.extract_units(data, "incidents.csv")
    assert len(units) == 1
    text = units[0][0]
    assert "incident: INC-1 | severity: P1 | owner: noc" in text
    assert "incident: INC-2 | severity: P3 | owner: support" in text
    # deterministic: same input, same output
    assert kb.extract_units(data, "incidents.csv") == units


def test_csv_header_only_and_empty():
    assert "a | b" in kb.extract_units(b"a,b\n", "h.csv")[0][0]
    assert kb.extract_units(b"", "empty.csv") == []


# ---- XLSX -------------------------------------------------------------------

def test_xlsx_multiple_sheets_labelled_by_sheet_name():
    data = _xlsx_bytes({
        "Incidents": [["id", "severity"], ["INC-1", "P1"]],
        "Contacts": [["team", "email"], ["noc", "noc@example.com"]],
    })
    units = kb.extract_units(data, "book.xlsx")
    labels = [label for _, label in units]
    assert labels == ["sheet Incidents", "sheet Contacts"]
    assert "INC-1 | P1" in units[0][0]
    assert "noc@example.com" in units[1][0]


def test_xlsx_empty_worksheet_is_skipped():
    data = _xlsx_bytes({"Empty": [], "Real": [["x"]]})
    units = kb.extract_units(data, "book.xlsx")
    assert [label for _, label in units] == ["sheet Real"]


def test_malformed_xlsx_fails_cleanly():
    with pytest.raises(ValueError, match="not a readable .xlsx"):
        kb.extract_units(b"not a zip archive at all", "broken.xlsx")


# ---- existing formats unchanged ---------------------------------------------

def test_txt_and_md_extraction_unchanged():
    assert kb.extract_units(b"plain text body", "notes.txt") == [("plain text body", "document")]
    assert kb.extract_units(b"# Heading\n\nbody", "notes.md") == [("# Heading\n\nbody", "document")]


def test_pdf_extraction_still_uses_pypdf_per_page():
    import pypdf
    assert hasattr(pypdf, "PdfReader")  # pypdf is still the PDF path
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    # a blank page yields no text, which must stay an empty unit list
    assert kb.extract_units(buffer.getvalue(), "blank.pdf") == []


def test_chunking_semantics_unchanged():
    units = [("a" * 2000, "document")]
    chunks = kb.chunk_units(units, size=800, overlap=120)
    assert len(chunks) == 3
    assert chunks[1][1]["location"] == "document, chars 680-1480"


# ---- upload path accepts the new types --------------------------------------

def test_upload_accepts_new_types_and_still_rejects_unknown(client):
    adapters.set_adapter_override(adapters.FakeModelAdapter())
    login(client, "admin@platform.local")

    r = client.post("/api/knowledge/upload",
                    files={"file": ("book.xlsx", io.BytesIO(_xlsx_bytes(
                        {"S1": [["router", "logs"], ["fiber", "links"]]})),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                    data={"name": "Ops Workbook"})
    assert r.status_code == 201, r.text
    source = r.json()
    assert source["status"] == "ready" and source["chunk_count"] > 0

    detail = client.get(f"/api/knowledge/{source['id']}").json()
    assert "sheet S1" in str(detail["chunk_preview"][0]["meta"])

    bad = client.post("/api/knowledge/upload",
                      files={"file": ("x.exe", io.BytesIO(b"MZ"), "application/octet-stream")},
                      data={"name": "Nope"})
    assert bad.status_code == 422  # allow-list still enforced


def test_malformed_upload_uses_existing_failure_path(client):
    """A broken file must fail the source honestly, not 500 the app."""
    adapters.set_adapter_override(adapters.FakeModelAdapter())
    login(client, "admin@platform.local")
    r = client.post("/api/knowledge/upload",
                    files={"file": ("broken.docx", io.BytesIO(b"not a docx"),
                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                    data={"name": "Broken Doc"})
    assert r.status_code == 201            # upload succeeded, ingestion did not
    assert r.json()["status"] == "failed"  # existing error behaviour
    assert "not a readable .docx" in r.json()["error"]


# ---- bounded embedding retry -------------------------------------------------

class _FlakyAdapter:
    """Fails `fail_times` with a transient error, then succeeds."""
    model_id = "fake:flaky"

    def __init__(self, fail_times: int, error=ModelCallError):
        self.fail_times = fail_times
        self.error = error
        self.calls = 0

    def embed(self, texts):
        self.calls += 1
        if self.calls <= self.fail_times:
            raise self.error("transient provider blip")
        return [[0.1] * 8 for _ in texts]


def test_retry_recovers_from_a_transient_failure(monkeypatch):
    monkeypatch.setattr(kb, "EMBED_BACKOFF_SECONDS", (0, 0))
    adapter = _FlakyAdapter(fail_times=1)
    result = kb.embed_with_retry(adapter, ["a", "b"])
    assert len(result) == 2
    assert adapter.calls == 2  # failed once, succeeded on the retry


def test_retry_budget_is_bounded(monkeypatch):
    monkeypatch.setattr(kb, "EMBED_BACKOFF_SECONDS", (0, 0))
    adapter = _FlakyAdapter(fail_times=99)
    with pytest.raises(ModelCallError):
        kb.embed_with_retry(adapter, ["a"])
    assert adapter.calls == kb.EMBED_MAX_ATTEMPTS == 3  # no infinite retry


def test_permanent_failure_is_not_retried(monkeypatch):
    """ModelUnavailable means unconfigured — retrying cannot help."""
    monkeypatch.setattr(kb, "EMBED_BACKOFF_SECONDS", (0, 0))
    adapter = _FlakyAdapter(fail_times=99, error=ModelUnavailable)
    with pytest.raises(ModelUnavailable):
        kb.embed_with_retry(adapter, ["a"])
    assert adapter.calls == 1


def test_exhausted_retries_reach_the_existing_degradation_path(client, monkeypatch):
    """The whole point: retry may DELAY the outcome, never change it. After the
    budget is spent the source is still stored, keyword-only, and honest."""
    monkeypatch.setattr(kb, "EMBED_BACKOFF_SECONDS", (0, 0))
    adapters.set_adapter_override(_FlakyAdapter(fail_times=99))
    login(client, "admin@platform.local")
    r = client.post("/api/knowledge/upload",
                    files={"file": ("runbook.txt", io.BytesIO(
                        b"Check the core router logs, then verify fiber links."), "text/plain")},
                    data={"name": "Retry Exhausted Source"})
    assert r.status_code == 201
    source = r.json()
    assert source["status"] == "ready"        # chunks retained
    assert source["embedded"] is False        # honest degradation
    assert "keyword_only" in source["retrieval_mode"]
    assert "embedding unavailable" in source["error"]


def test_retry_succeeds_end_to_end_and_embeds(client, monkeypatch):
    monkeypatch.setattr(kb, "EMBED_BACKOFF_SECONDS", (0, 0))
    adapters.set_adapter_override(_FlakyAdapter(fail_times=1))
    login(client, "admin@platform.local")
    r = client.post("/api/knowledge/upload",
                    files={"file": ("runbook.txt", io.BytesIO(b"router logs and fiber links"),
                                    "text/plain")},
                    data={"name": "Retry Recovered Source"})
    assert r.status_code == 201
    assert r.json()["embedded"] is True   # recovered on retry
    assert r.json()["error"] is None
